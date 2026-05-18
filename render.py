import base64
import io
import tomllib
import subprocess
import urllib.request
from pathlib import Path
from jinja2 import Environment, FileSystemLoader
from weasyprint import HTML, CSS
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path(__file__).parent
TEMPLATES = BASE / "templates"
OUTPUT = BASE / "output"

LABELS = {
    "da": {
        "contact": "Kontakt",
        "languages": "Sprog",
        "skills": "Kompetencer",
        "about": "Om mig",
        "experience": "Joberfaring",
        "education": "Uddannelse",
        "profile": "Profil",
    },
    "en": {
        "contact": "Contact",
        "languages": "Languages",
        "skills": "Skills",
        "about": "About me",
        "experience": "Work experience",
        "education": "Education",
        "profile": "Profile",
    },
}


IPFS_GW = "http://localhost:8080"


def strip_meta(cv: dict) -> dict:
    return {k: v for k, v in cv.items() if not k.startswith("_")}


def _is_cid(photo: str) -> bool:
    return bool(photo) and "/" not in photo


def _fetch_photo_bytes(photo: str) -> bytes | None:
    """Hent billedbytes fra IPFS-gateway eller lokalt filsystem (bagudkompatibilitet)."""
    if not photo:
        return None
    if not _is_cid(photo):
        p = BASE / photo
        return p.read_bytes() if p.exists() else None
    try:
        with urllib.request.urlopen(f"{IPFS_GW}/ipfs/{photo}", timeout=15) as r:
            return r.read()
    except Exception:
        return None


def _photo_data_url(photo: str) -> str:
    data = _fetch_photo_bytes(photo)
    if not data:
        return ""
    return "data:image/jpeg;base64," + base64.b64encode(data).decode()


def load_cv(path: Path | None = None) -> dict:
    with open(path or BASE / "cv.toml", "rb") as f:
        return tomllib.load(f)


def render_html(cv: dict) -> str:
    cv = strip_meta(cv)
    cv["personal"] = dict(cv.get("personal", {}))
    photo = cv["personal"].get("photo", "")
    if photo:
        cv["personal"]["photo"] = _photo_data_url(photo)
    lang = cv.get("lang", "da")
    labels = LABELS.get(lang, LABELS["da"])
    env = Environment(loader=FileSystemLoader(str(TEMPLATES)))
    tmpl = env.get_template("cv.html")
    return tmpl.render(cv=cv, labels=labels)


def to_pdf(cv: dict, out_path: Path | None = None) -> Path:
    html_str = render_html(cv)
    css_path = TEMPLATES / "cv.css"
    out = out_path or OUTPUT / "cv.pdf"
    out.parent.mkdir(parents=True, exist_ok=True)
    HTML(string=html_str, base_url=str(BASE)).write_pdf(
        str(out),
        stylesheets=[CSS(filename=str(css_path))],
    )
    return out


def _add_run(para, text, bold=False, italic=False, size=None, color=None):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def _set_para_spacing(para, before=0, after=0):
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)


def to_docx(cv: dict, out_path: Path | None = None) -> Path:
    cv = strip_meta(cv)
    lang = cv.get("lang", "da")
    labels = LABELS.get(lang, LABELS["da"])
    doc = Document()
    out = out_path or OUTPUT / "cv.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    p = cv["personal"]
    name_para = doc.add_paragraph()
    _set_para_spacing(name_para, after=2)
    _add_run(name_para, p["name"], bold=True, size=22)

    title_para = doc.add_paragraph()
    _set_para_spacing(title_para, after=4)
    _add_run(title_para, p["title"], size=14, color=(70, 130, 180))

    contact = doc.add_paragraph()
    _set_para_spacing(contact, after=8)
    _add_run(contact, f"{p['email']}  |  {p['phone']}  |  {p['address']}", size=10)

    if p.get("photo"):
        photo_bytes = _fetch_photo_bytes(p["photo"])
        if photo_bytes:
            photo_para = doc.add_paragraph()
            photo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_para_spacing(photo_para, after=6)
            run = photo_para.add_run()
            run.add_picture(io.BytesIO(photo_bytes), width=Cm(2.5))

    doc.add_heading(labels["profile"], level=2)
    s = doc.add_paragraph(cv["summary"]["text"].strip())
    _set_para_spacing(s, after=8)

    doc.add_heading(labels["skills"], level=2)
    skills_para = doc.add_paragraph(", ".join(cv["skills"]["tags"]))
    _set_para_spacing(skills_para, after=8)

    doc.add_heading(labels["languages"], level=2)
    for lang_item in cv["languages"]:
        dots = "●" * lang_item["level"] + "○" * (4 - lang_item["level"])
        lp = doc.add_paragraph()
        _set_para_spacing(lp, after=2)
        _add_run(lp, f"{lang_item['name']}  {dots}", size=10)

    doc.add_heading(labels["experience"], level=1)
    for job in cv["experience"]:
        h = doc.add_paragraph()
        _set_para_spacing(h, before=6, after=0)
        _add_run(h, job["title"], bold=True, size=12)
        _add_run(h, f"  —  {job['company']} ({job['duration']})", size=10, color=(100, 100, 100))

        period_p = doc.add_paragraph()
        _set_para_spacing(period_p, after=2)
        _add_run(period_p, job["period"], italic=True, size=9, color=(130, 130, 130))

        if job.get("type"):
            tp = doc.add_paragraph()
            _set_para_spacing(tp, after=2)
            _add_run(tp, job["type"], size=10)

        if job.get("description"):
            dp = doc.add_paragraph(job["description"])
            _set_para_spacing(dp, after=2)

        for bullet in job.get("bullets", []):
            bp = doc.add_paragraph(bullet, style="List Bullet")
            _set_para_spacing(bp, after=1)

        for sub in job.get("subsections", []):
            sp = doc.add_paragraph()
            _set_para_spacing(sp, before=4, after=1)
            _add_run(sp, sub["heading"], bold=True, size=10)
            for bullet in sub["bullets"]:
                bp = doc.add_paragraph(bullet, style="List Bullet")
                _set_para_spacing(bp, after=1)

    doc.add_heading(labels["education"], level=1)
    for edu in cv["education"]:
        h = doc.add_paragraph()
        _set_para_spacing(h, before=6, after=0)
        _add_run(h, edu["degree"], bold=True, size=12)
        _add_run(h, f"  —  {edu['institution']} ({edu['duration']})", size=10, color=(100, 100, 100))

        period_p = doc.add_paragraph()
        _set_para_spacing(period_p, after=2)
        _add_run(period_p, edu["period"], italic=True, size=9, color=(130, 130, 130))

        for bullet in edu.get("bullets", []):
            bp = doc.add_paragraph(bullet, style="List Bullet")
            _set_para_spacing(bp, after=1)

    doc.save(str(out))
    return out


def to_odf(cv: dict) -> Path:
    docx_path = to_docx(cv)
    out = OUTPUT / "cv.odt"
    subprocess.run([
        "libreoffice", "--headless", "--convert-to", "odt",
        "--outdir", str(OUTPUT), str(docx_path),
    ], check=True, capture_output=True)
    return out
